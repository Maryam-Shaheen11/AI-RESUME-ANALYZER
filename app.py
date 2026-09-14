import os
import json
import re
from io import BytesIO

import streamlit as st
from groq import Groq
from pypdf import PdfReader
from docx import Document


# -----------------------------
# Page configuration
# -----------------------------
st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide",
)

st.title("📄 AI Resume Analyzer")
st.caption("Upload your resume, paste a job description, and get an AI-powered match analysis.")


# -----------------------------
# Resume text extraction
# -----------------------------
def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from a PDF file."""
    reader = PdfReader(BytesIO(file_bytes))
    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract text from a DOCX file."""
    document = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also read text from tables, which are common in resumes.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()


def extract_txt_text(file_bytes: bytes) -> str:
    """Extract text from a TXT file."""
    return file_bytes.decode("utf-8", errors="ignore").strip()


def extract_resume_text(uploaded_file) -> str:
    """Extract resume text based on file type."""
    file_bytes = uploaded_file.getvalue()
    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_pdf_text(file_bytes)

    if file_name.endswith(".docx"):
        return extract_docx_text(file_bytes)

    if file_name.endswith(".txt"):
        return extract_txt_text(file_bytes)

    raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")


# -----------------------------
# AI analysis
# -----------------------------
def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict:
    """Send resume + job description to the Groq-hosted model and return JSON."""
    client = Groq(api_key=api_key)

    system_prompt = """
You are an expert ATS resume analyzer and career coach.

Analyze the candidate resume against the supplied job description.

Return ONLY valid JSON. Do not use Markdown fences.

Use exactly this JSON structure:
{
  "match_score": 0,
  "overall_summary": "",
  "matching_skills": [],
  "missing_skills": [],
  "matching_keywords": [],
  "missing_keywords": [],
  "experience_match": "",
  "education_match": "",
  "ats_issues": [],
  "resume_improvements": [],
  "recommended_changes": [],
  "final_verdict": ""
}

Rules:
- match_score must be an integer from 0 to 100.
- Be evidence-based. Do not invent experience, education, certifications, projects, or skills.
- Distinguish between skills explicitly present in the resume and skills only implied.
- For missing skills, list only important job requirements that are not clearly supported by the resume.
- Identify ATS problems such as missing keywords, weak section wording, formatting concerns, lack of measurable achievements, or irrelevant content.
- Keep recommendations practical and specific.
- If the resume does not contain enough information to judge something, say so.
"""

    user_prompt = f"""
RESUME:
{resume_text}

JOB DESCRIPTION:
{job_description}
"""

    response = client.chat.completions.create(
        model="openai/gpt-oss-120b",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.2,
        max_completion_tokens=3000,
    )

    content = response.choices[0].message.content.strip()

    # Remove accidental Markdown fences if the model adds them.
    content = re.sub(r"^```(?:json)?\s*", "", content, flags=re.IGNORECASE)
    content = re.sub(r"\s*```$", "", content)

    try:
        result = json.loads(content)
    except json.JSONDecodeError as exc:
        raise ValueError(
            "The AI returned an invalid response. Please try the analysis again."
        ) from exc

    # Basic validation / safe defaults.
    result["match_score"] = max(0, min(100, int(result.get("match_score", 0))))
    list_fields = [
        "matching_skills",
        "missing_skills",
        "matching_keywords",
        "missing_keywords",
        "ats_issues",
        "resume_improvements",
        "recommended_changes",
    ]

    for field in list_fields:
        if not isinstance(result.get(field), list):
            result[field] = []

    for field in ["overall_summary", "experience_match", "education_match", "final_verdict"]:
        if not isinstance(result.get(field), str):
            result[field] = ""

    return result


# -----------------------------
# Sidebar
# -----------------------------
with st.sidebar:
    st.header("⚙️ Settings")
    st.info(
        "This app uses the Groq API and the hosted "
        "`openai/gpt-oss-120b` model."
    )

    api_key = st.text_input(
        "Groq API Key",
        type="password",
        value=os.getenv("GROQ_API_KEY", ""),
        help="You can also set GROQ_API_KEY as an environment variable.",
    )

    st.markdown("---")
    st.write("**Supported resume formats:**")
    st.write("• PDF")
    st.write("• DOCX")
    st.write("• TXT")


# -----------------------------
# Main interface
# -----------------------------
col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Upload Resume")
    uploaded_file = st.file_uploader(
        "Choose your resume",
        type=["pdf", "docx", "txt"],
        help="For best results, use a text-based PDF or DOCX.",
    )

with col2:
    st.subheader("2. Job Description")
    job_description = st.text_area(
        "Paste the complete job description here",
        height=260,
        placeholder="Example: We are looking for a Python/AI Engineer with experience in machine learning, LLMs, APIs...",
    )

if uploaded_file:
    try:
        resume_text = extract_resume_text(uploaded_file)

        if resume_text:
            st.success(
                f"Resume loaded successfully: {uploaded_file.name} "
                f"({len(resume_text):,} characters extracted)"
            )

            with st.expander("👀 Preview extracted resume text"):
                st.text(resume_text[:10000])
        else:
            resume_text = ""
            st.error(
                "No readable text was found in this file. "
                "If it is a scanned/image-only PDF, please use a text-based PDF or DOCX."
            )
    except Exception as exc:
        resume_text = ""
        st.error(f"Could not read the resume: {exc}")
else:
    resume_text = ""


st.markdown("---")

analyze_clicked = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_clicked:
    if not api_key.strip():
        st.error("Please enter your Groq API key in the sidebar.")
        st.stop()

    if not uploaded_file:
        st.error("Please upload a resume first.")
        st.stop()

    if not resume_text.strip():
        st.error("The uploaded resume does not contain readable text.")
        st.stop()

    if len(job_description.strip()) < 50:
        st.error("Please enter a more complete job description.")
        st.stop()

    with st.spinner("Analyzing your resume against the job description..."):
        try:
            analysis = analyze_resume(
                resume_text=resume_text,
                job_description=job_description,
                api_key=api_key.strip(),
            )
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            st.stop()

    st.success("Analysis completed successfully! 🎉")

    # Score
    score = analysis["match_score"]

    st.subheader("📊 Resume Match Score")
    st.progress(score / 100)
    st.metric("Overall Match", f"{score}%")

    # Summary
    st.subheader("📝 Overall Summary")
    st.write(analysis["overall_summary"])

    # Skills
    left, right = st.columns(2)

    with left:
        st.subheader("✅ Matching Skills")
        if analysis["matching_skills"]:
            for item in analysis["matching_skills"]:
                st.write(f"• {item}")
        else:
            st.write("No strong matching skills were identified.")

    with right:
        st.subheader("❌ Missing / Weak Skills")
        if analysis["missing_skills"]:
            for item in analysis["missing_skills"]:
                st.write(f"• {item}")
        else:
            st.write("No major missing skills were identified.")

    # Keywords
    left, right = st.columns(2)

    with left:
        st.subheader("🔑 Matching Keywords")
        if analysis["matching_keywords"]:
            st.write(", ".join(str(x) for x in analysis["matching_keywords"]))
        else:
            st.write("No matching keywords identified.")

    with right:
        st.subheader("⚠️ Missing Keywords")
        if analysis["missing_keywords"]:
            st.write(", ".join(str(x) for x in analysis["missing_keywords"]))
        else:
            st.write("No important missing keywords identified.")

    # Detailed checks
    st.subheader("💼 Experience Match")
    st.write(analysis["experience_match"])

    st.subheader("🎓 Education Match")
    st.write(analysis["education_match"])

    st.subheader("🤖 ATS Issues")
    if analysis["ats_issues"]:
        for item in analysis["ats_issues"]:
            st.write(f"• {item}")
    else:
        st.write("No major ATS issues identified.")

    st.subheader("✨ Resume Improvements")
    if analysis["resume_improvements"]:
        for item in analysis["resume_improvements"]:
            st.write(f"• {item}")
    else:
        st.write("No major improvement suggestions.")

    st.subheader("🚀 Recommended Changes")
    if analysis["recommended_changes"]:
        for item in analysis["recommended_changes"]:
            st.write(f"• {item}")
    else:
        st.write("No additional changes recommended.")

    st.subheader("🎯 Final Verdict")
    st.info(analysis["final_verdict"])

    # Download JSON report
    report_json = json.dumps(analysis, indent=2, ensure_ascii=False)

    st.download_button(
        "⬇️ Download Analysis Report (JSON)",
        data=report_json,
        file_name="resume_analysis_report.json",
        mime="application/json",
    )

st.markdown("---")
st.caption("AI Resume Analyzer • Built with Python + Streamlit + Groq")
